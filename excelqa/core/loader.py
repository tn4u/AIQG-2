"""Đọc Excel lộn xộn -> danh sách Table sạch (KHÔNG dùng LLM).

Giải quyết: (1) header nhiều tầng + ô gộp + tiêu đề bị ngắt dòng, (2) nhãn hàng lặp qua các khối
(`section`), (3) hàng tổng/cha/con, (4) số lưu dạng chuỗi ("$2,253", "10%", "(94)", "..", "(-8-)"),
(5) cột nhãn không nằm ở cột A, (6) nhiều sheet / nhiều bảng trên một sheet.
"""
from __future__ import annotations
import datetime as dt
import re
from collections import Counter
from dataclasses import dataclass, field

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.utils import column_index_from_string, get_column_letter

NA_RE = re.compile(r"^(\$?\s*[–—-]+|\.{2,}|…|\(-?\d+-?\)|\(x\)|x|n/?a|\(na\)|\*+|nan)$", re.I)
KEEP_HYPHEN = {"self", "non", "semi", "full", "part", "year", "high", "low", "co"}
FOOT_RE = re.compile(r"^\s*(\(\d+\)|\d+/|footnotes?\b|notes?\b|source\b|contact\b|\*)", re.I)
GROUP_RE = re.compile(r"^(?P<g>[^\d,][^,]*?),\s*\d+\s+\D[^,)]*$")
TOTAL_RE = re.compile(r"^(total\b|all items\b|all industries\b)", re.I)
HELPERS = ["_row", "section", "label", "is_total", "level", "has_children"]


@dataclass
class Table:
    file: str
    sheet: str
    name: str
    title: str
    notes: list
    year: str | None
    label_col: str
    col_names: list          # tên đầy đủ các cột giá trị (không gồm cột nhãn)
    col_map: dict            # c1 -> tên đầy đủ
    df: pd.DataFrame
    footnotes: list = field(default_factory=list)


# ---------------------------------------------------------------- tiện ích
def clean_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float, np.integer, np.floating)):
        if isinstance(v, float) and np.isnan(v):
            return ""
        return str(int(v)) if float(v).is_integer() else str(v)
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).replace("\xa0", " ")

    def _join(m):  # "Con-\nstruction" -> "Construction"; "Self-\nemployed" giữ gạch nối
        keep = "-" if m.group(1).lower() in KEEP_HYPHEN else ""
        return m.group(1) + keep + m.group(2)

    s = re.sub(r"([A-Za-z]+)-\s*\n\s*([a-z])", _join, s)
    return re.sub(r"\s+", " ", s).strip()


def strip_footnote(s: str, label: bool = False) -> str:
    s = re.sub(r"(\s*\(\(?[a-z0-9]{1,2}\)\)?)+$", "", s)   # "(1)(3)", "(a)", "((a)(f)"
    s = re.sub(r"\s+\d+/$", "", s)                          # "Trade balance 3/"
    if label:
        s = re.sub(r"(?<=[a-z\)])\d(?:,\d)*\s*$", "", s)    # "Private industry3" (không đụng "Q2")
    return s.strip()


def is_num(v) -> bool:
    return isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, bool) \
        and not (isinstance(v, float) and np.isnan(v))


def is_date(v) -> bool:
    return isinstance(v, (dt.datetime, dt.date))


def is_na(v) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and np.isnan(v):
        return True
    return isinstance(v, str) and (v.strip() == "" or bool(NA_RE.match(v.strip())))


def parse_num(v):
    """Số thật hoặc chuỗi số kiểu '$2,253', '10%', '(94)', '-0.5'. Không phải số -> None."""
    if is_num(v):
        return v
    if isinstance(v, str):
        s = v.strip()
        neg = s.startswith("(") and s.endswith(")")
        s = re.sub(r"[\$,\s%]", "", s.strip("()"))
        if re.fullmatch(r"-?\d+(\.\d+)?", s):
            f = float(s)
            f = -abs(f) if neg else f
            return int(f) if "." not in s else f
    return None


def is_numlike(v) -> bool:
    return is_num(v) or is_date(v) or (isinstance(v, str) and parse_num(v) is not None)


def is_year(v) -> bool:
    return is_num(v) and float(v).is_integer() and 1900 <= v <= 2100


def is_yearlike(v) -> bool:
    return is_year(v) or (isinstance(v, str) and bool(re.fullmatch(r"(19|20)\d{2}", strip_footnote(clean_text(v)))))


# ---------------------------------------------------------------- đọc grid
def _read_grid(ws_v, ws_f):
    nr, nc = ws_v.max_row, ws_v.max_column
    grid = [[ws_v.cell(r, c).value for c in range(1, nc + 1)] for r in range(1, nr + 1)]
    for r in range(nr):  # ô công thức không có giá trị cache: giải "=H1" đơn giản
        for c in range(nc):
            if grid[r][c] is None:
                f = ws_f.cell(r + 1, c + 1).value
                if isinstance(f, str) and f.startswith("="):
                    m = re.fullmatch(r"=\$?([A-Z]+)\$?(\d+)", f.strip())
                    if m:
                        grid[r][c] = ws_v.cell(int(m.group(2)), column_index_from_string(m.group(1))).value
    spans = {}
    for rng in ws_v.merged_cells.ranges:  # giải ô gộp: điền giá trị anchor vào cả vùng
        v = grid[rng.min_row - 1][rng.min_col - 1]
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                grid[r - 1][c - 1] = v
        if rng.min_row == rng.max_row:
            spans.setdefault(rng.min_row - 1, []).append((rng.min_col - 1, rng.max_col - 1))
    return grid, spans


def _pick_label_col(grid, used, fw):
    """Cột nhãn = cột 'dày' trái nhất chủ yếu là chữ (hoặc năm). Cột thưa bên trái = tiêu đề khối."""
    rows = [r for r in range(len(grid)) if r not in fw]
    stats = {}
    for c in used:
        vals = [grid[r][c] for r in rows if clean_text(grid[r][c]) != ""]
        stats[c] = (len(vals), sum(isinstance(v, str) for v in vals), sum(is_year(v) for v in vals))
    mx = max(n for n, _, _ in stats.values())
    for c in used:
        n, ns, ny = stats[c]
        if n >= 3 and n >= 0.5 * mx and (ns >= 0.5 * n or ny >= 0.8 * n):
            return c
    for c in used:
        n, ns, _ = stats[c]
        if n >= 3 and ns >= 0.5 * n:
            return c
    return used[0]


# ---------------------------------------------------------------- hàm chính
def load_tables(path, file_name: str | None = None) -> list[Table]:
    wb_v = openpyxl.load_workbook(path, data_only=True, keep_links=False)
    wb_f = openpyxl.load_workbook(path, keep_links=False)
    out = []
    try:
        for ws_v in wb_v.worksheets:
            try:
                out += _load_sheet(ws_v, wb_f[ws_v.title], file_name or str(path))
            except Exception:
                continue          # một sheet lỗi không làm mất cả file
    finally:
        wb_v.close(); wb_f.close()
    return out


def _load_sheet(ws_v, ws_f, fname) -> list[Table]:
    grid, spans = _read_grid(ws_v, ws_f)
    nr = len(grid)
    used = [c for c in range(ws_v.max_column) if any(clean_text(grid[r][c]) != "" for r in range(nr))]
    if not used:
        return []
    fw = {r for r, sp in spans.items() if len(used) > 1 and any(a <= used[0] and b >= used[-1] for a, b in sp)}
    lc = _pick_label_col(grid, used, fw)
    # gộp ngang phủ từ cột nhãn đến hết bảng (vd tiêu đề B1:G1 khi cột A là cột phụ) cũng là tiêu đề/ghi chú
    fw |= {r for r, sp in spans.items() if len(used) > 1 and any(a <= lc and b >= used[-1] for a, b in sp)}
    cnts = {c: sum(1 for r in range(nr) if r not in fw and clean_text(grid[r][c]) != "") for c in used}
    mx = max(cnts.values())
    left = [c for c in used if c < lc and cnts[c] < 0.5 * mx]   # cột thưa bên trái = tiêu đề khối, không phải dữ liệu
    used = [c for c in used if c not in left]
    dcols = [c for c in used if c != lc]

    def row_label(r):
        lab = clean_text(grid[r][lc])
        if lab == "":                       # tiêu đề khối nằm ở cột thưa bên trái
            for c in left:
                t = clean_text(grid[r][c])
                if t and not re.fullmatch(r"-?\d+", t):
                    return t
        return lab

    def classify(r):
        """-> (kind, số ô không rỗng)"""
        if r in fw:
            return "fullwidth", 0
        lab = row_label(r)
        others = [grid[r][c] for c in dcols]
        ne = [v for v in others if not is_na(v)]
        has_na = any(isinstance(v, str) and v.strip() != "" and is_na(v) for v in others)
        if not ne:
            if has_na and lab:
                return "data", 0            # hàng toàn dấu '–' vẫn là dữ liệu
            if lab == "":
                return "empty", 0
            return ("footnote" if FOOT_RE.match(lab) else "section"), 0
        if all(is_yearlike(v) for v in ne):
            return "header", len(ne)
        if sum(is_numlike(v) for v in ne) >= 0.5 * len(ne):
            return "data", len(ne)
        return "header", len(ne)

    title_lines, footnotes, tables = [], [], []
    state, cur, stack, last_kind, last_yq = "pre", None, [], None, None

    def new_table():
        return {"header": [], "data": []}

    def push_section(label):
        nonlocal stack
        if label.isupper() and label != "TOTAL":
            stack = [label]
        elif stack and stack[0].isupper() and stack[0] != "TOTAL":
            stack = [stack[0], label]
        else:
            stack = [label]

    for r in range(nr):
        k, n_ne = classify(r)
        raw = grid[r][lc] if clean_text(grid[r][lc]) != "" else None
        lab = strip_footnote(row_label(r), label=True)
        if k == "empty":
            continue
        if k == "fullwidth":
            (title_lines if state == "pre" else footnotes).append(clean_text(grid[r][used[0]]))
            continue
        if k == "footnote":
            footnotes.append(clean_text(grid[r][lc]))
            continue
        if k == "header":
            if state == "data":
                if last_kind == "section":      # dòng mã/ghi chú ngay sau tiêu đề khối: bỏ qua
                    last_kind = "header"
                    continue
                if n_ne < 0.5 * len(dcols):     # ô chữ lẻ giữa vùng dữ liệu: coi là ghi chú
                    footnotes.append(" ".join(clean_text(v) for v in (grid[r][c] for c in dcols) if clean_text(v)))
                    continue
                tables.append(cur)
                cur, stack = new_table(), []
            elif cur is None:
                cur = new_table()
            cur["header"].append(r)
            state, last_kind = "head", "header"
            continue
        if k == "section":
            last_kind = "section"
            if state == "pre":
                title_lines.append(lab)
                continue
            state = "data"
            push_section(lab)
            continue
        # k == "data"
        last_kind = "data"
        if cur is None:
            cur = new_table()
        state = "data"
        m = re.fullmatch(r"(\d{4}) Q[1-4]", lab)
        if m:
            last_yq = m.group(1)
        elif re.fullmatch(r"Q[1-4]", lab) and last_yq:   # "Q2" -> "1979 Q2"
            lab = f"{last_yq} {lab}"
        is_group = bool(GROUP_RE.match(lab))
        if is_group:
            push_section(lab)
            section = " > ".join(stack[:-1])
        else:
            section = " > ".join(stack)
        indent = ws_f.cell(r + 1, lc + 1).alignment.indent or 0
        lead = (len(raw) - len(raw.replace("\xa0", " ").lstrip())) if isinstance(raw, str) else 0
        cur["data"].append({"r": r, "label": lab, "section": section, "indent": int(indent), "lead": lead,
                            "is_total": bool(TOTAL_RE.match(lab)) or is_group})
    if cur is not None and (cur["data"] or cur["header"]):
        tables.append(cur)

    title = " / ".join(t for t in title_lines if t)
    hmerge = {(r, c) for r, sp in spans.items() for a, b in sp if b > a for c in range(a, b + 1)}
    result = []
    for t in tables:
        if t["data"]:
            result.append(_build_table(t, grid, used, lc, dcols, ws_v.title, fname, title, hmerge))
    if result:
        result[-1].footnotes = [f for f in footnotes if f]
    return result


def _unique(names):
    seen, out = {}, []
    for n in names:
        if n in seen or n in HELPERS:
            seen[n] = seen.get(n, 1) + 1
            n = f"{n} #{seen[n]}"
        seen.setdefault(n, 1)
        out.append(n)
    return out


def _wraps(prev: str, cur: str) -> bool:
    return prev.endswith((",", "&", "-")) or cur[:1].islower() or cur[:1] == "&"


def _join_parts(parts) -> str:
    """parts = [(text, không_nằm_trong_ô_gộp)]. Mảnh bị ngắt dòng của một nhãn -> nối bằng dấu cách;
    tầng phân cấp -> ' | '."""
    out = []
    for text, uniq in parts:
        if out and out[-1][1] and uniq and _wraps(out[-1][0], text):
            prev = out[-1][0]
            if prev.endswith("-") and text[:1].islower():
                out[-1] = (prev[:-1] + text, True)
            else:
                out[-1] = (prev + " " + text, True)
        else:
            out.append((text, uniq))
    return " | ".join(t for t, _ in out)


def _build_table(t, grid, used, lc, dcols, sheet, fname, title, hmerge=frozenset()) -> Table:
    hdr = t["header"]
    year, notes, per_rows = None, [], []
    for r in hdr:
        vals = [strip_footnote(clean_text(grid[r][c])) for c in dcols]
        ne = [v for v in vals if v]
        if len(dcols) >= 2 and len(ne) == len(dcols) and len(set(ne)) == 1:
            if re.fullmatch(r"\d{4}", ne[0]):
                year = ne[0]
            else:
                notes.append(ne[0])
        elif len(dcols) >= 3 and len(ne) == 1:     # ghi chú lẻ loi trên header (vd 'Seasonally adjusted')
            notes.append(ne[0])
        else:
            per_rows.append(r)
    names = []
    for c in used:
        parts = []
        for r in (hdr if c == lc else per_rows):
            v = strip_footnote(clean_text(grid[r][c]))
            if v and (not parts or parts[-1][0] != v):
                parts.append((v, (r, c) not in hmerge))
        names.append(_join_parts(parts) or f"col_{get_column_letter(c + 1)}")
    names = _unique(names)
    name_of = dict(zip(used, names))
    tname = sheet
    if per_rows:
        for c in dcols:
            v = strip_footnote(clean_text(grid[per_rows[0]][c]))
            if v:
                tname = v
                break

    recs = []
    for e in t["data"]:
        rec = {"_row": e["r"] + 1, "section": e["section"], "label": e["label"],
               "is_total": e["is_total"], "level": e["indent"], "has_children": False}
        for c in used:
            rec[name_of[c]] = e["label"] if c == lc else grid[e["r"]][c]
        recs.append(rec)
    df = pd.DataFrame(recs)
    for c in dcols:
        n = name_of[c]
        if re.search(r"\b(code|naics)\b", n, re.I):      # mã (NAICS...) giữ dạng chữ để đếm độ dài
            df[n] = pd.Series([None if is_na(v) else clean_text(v) for v in df[n]], index=df.index, dtype="object")
        else:
            df[n] = _coerce(df[n])

    ind = [n for n in names if n.lower().startswith("indent") and pd.api.types.is_numeric_dtype(df[n])]
    if ind:
        df["level"] = df[ind[0]].fillna(0).astype(int)
    elif df["level"].max() == 0:
        leads = sorted({e["lead"] for e in t["data"]})
        if len(leads) > 1:                                # thụt lề bằng dấu cách đầu nhãn
            rank = {v: i for i, v in enumerate(leads)}
            df["level"] = [rank[e["lead"]] for e in t["data"]]
        else:
            _infer_hierarchy(df, [name_of[c] for c in dcols])
    lv = df["level"].tolist(); sec = df["section"].tolist()
    df["has_children"] = [i + 1 < len(df) and lv[i + 1] > lv[i] and sec[i + 1] == sec[i] for i in range(len(df))]

    value_names = [name_of[c] for c in dcols]
    return Table(fname, sheet, tname, title, notes, year, name_of[lc], value_names,
                 {f"c{i + 1}": n for i, n in enumerate(value_names)}, df)


def _coerce(s: pd.Series) -> pd.Series:
    vals = [None if is_na(v) else (parse_num(v) if isinstance(v, str) and parse_num(v) is not None else v) for v in s]
    nn = [v for v in vals if v is not None]
    if not nn:
        return pd.Series([np.nan] * len(vals), index=s.index, dtype="float64")
    n_date = sum(is_date(v) for v in nn)
    n_num = sum(is_num(v) for v in nn)
    if n_date >= 0.6 * len(nn):
        return pd.to_datetime(pd.Series([v if is_date(v) else None for v in vals], index=s.index))
    if n_num >= 0.5 * len(nn):
        num = pd.Series([v if is_num(v) else np.nan for v in vals], index=s.index, dtype="float64")
        if num.notna().all() and (num % 1 == 0).all():
            return num.astype("int64")
        return num
    return pd.Series([None if v is None else clean_text(v) for v in vals], index=s.index, dtype="object")


def _infer_hierarchy(df: pd.DataFrame, value_cols: list, tol: float = 0.05):
    """Không có thụt lề: suy cha/con khi giá trị hàng i ≈ tổng k hàng liền sau (k>=2)."""
    col = next((c for c in value_cols if pd.api.types.is_numeric_dtype(df[c]) and df[c].notna().mean() > 0.8), None)
    if col is None:
        return
    v = df[col].tolist(); sec = df["section"].tolist(); lv = [0] * len(df)
    for i in range(len(df)):
        if not df["is_total"].iloc[i] and not np.isnan(v[i]) and v[i] > 0:
            acc = 0.0
            for k in range(1, 9):
                j = i + k
                if j >= len(df) or sec[j] != sec[i] or np.isnan(v[j]) or v[j] < 0 or v[j] > v[i] + tol:
                    break
                acc += v[j]
                if k >= 2 and abs(acc - v[i]) <= tol:
                    for m in range(i + 1, j + 1):
                        lv[m] = max(lv[m], lv[i] + 1)
                    break
    df["level"] = lv